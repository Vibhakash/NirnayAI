from typing import Tuple, List, Dict
from docx import Document


def extract_text_from_docx(file_path: str) -> Tuple[str, List[Dict]]:
    """
    Extract text from a .docx file.
    Returns (full_text, sections) where sections represent paragraphs/tables.
    """
    doc = Document(file_path)
    sections = []
    full_parts = []
    offset = 0

    # Paragraphs
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text:
            sections.append({
                "section_type": "paragraph",
                "index": i,
                "text": text,
                "char_offset_start": offset,
                "char_offset_end": offset + len(text),
            })
            full_parts.append(text)
            offset += len(text) + 1  # +1 for newline

    # Tables
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                sections.append({
                    "section_type": "table_row",
                    "index": f"t{t_idx}_r{r_idx}",
                    "text": row_text,
                    "char_offset_start": offset,
                    "char_offset_end": offset + len(row_text),
                })
                full_parts.append(row_text)
                offset += len(row_text) + 1

    return "\n".join(full_parts), sections
